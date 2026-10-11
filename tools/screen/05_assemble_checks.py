"""Step 5: assemble xlsx/csv and run checks. Usage: 05_assemble_checks.py [--test]"""
import csv, statistics, sys
from collections import Counter
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from _util import *

TEST = "--test" in sys.argv
sfx = "_test" if TEST else ""
CAPS = json.loads((OUT / f"caps{sfx}.json").read_text())
NOP = {o["cik"]: o for o in json.loads((OUT / f"nopat{sfx}.json").read_text())}
DESC = json.loads((OUT / f"desc{sfx}.json").read_text())
today = date.fromisoformat(TODAY)
TARGET = 500


def shell(cik):
    p = RAW / f"cf_{cik:010d}.json"
    if not p.exists():
        return False
    best = None
    for arr in json.loads(p.read_text()).get("facts", {}).get("dei", {}).get("EntityShellCompany", {}).get("units", {}).values():
        for e in arr:
            if best is None or e["end"] > best["end"]:
                best = e
    return bool(best and best["val"] is True)


rows = []
for r in CAPS["rows"]:
    n, dsc = NOP[r["cik"]], DESC.get(r["ticker"], {})
    flags = list(r["flags"])
    if r["shares_src"] == "yahoo": flags.append("shares_from_yahoo")
    if n["nopat"] is None: flags.append("nopat_unavailable")
    elif n["end"] < "2025-07-11": flags.append("fy_older_than_15m")
    if n["src"] == "companyfacts": flags.append("nopat_from_companyfacts")
    if dsc.get("src") == "sic": flags.append("desc_sic_fallback")
    if shell(r["cik"]): flags.append("shell_company")
    m = lambda x: None if x is None else round(x / 1e6, 4)
    rows.append(dict(
        cik=r["cik"], ticker=r["ticker"], company=r["sec_name"] or r["name"], exchange=r["exchange"],
        market_cap_usd_m=round(r["market_cap"] / 1e6, 4), price=r["price"], price_date=r["price_date"],
        shares_m=round(r["shares"] / 1e6, 4), shares_date=r["shares_date"], sic_desc=r["sic_desc"],
        description=dsc.get("description"), fiscal_year=n["fy"], operating_income_m=m(n["op"]),
        tax_rate=None if n["tax_rate"] is None else round(n["tax_rate"], 4), nopat_m=m(n["nopat"]),
        nopat_sentence=n["sentence"], flags=";".join(flags)))
rows = rows[:TARGET]
for i, r in enumerate(rows, 1):
    r["rank"] = i
COLS = ["rank", "ticker", "company", "exchange", "market_cap_usd_m", "price", "price_date", "shares_m", "shares_date",
        "sic_desc", "description", "fiscal_year", "operating_income_m", "tax_rate", "nopat_m", "nopat_sentence", "flags"]
name = f"microcaps_500{sfx}"
with open(OUT / f"{name}.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=COLS, extrasaction="ignore"); w.writeheader(); w.writerows(rows)
wb = Workbook(); ws = wb.active; ws.title = "microcaps"
ws.append(COLS)
for c in ws[1]: c.font = Font(bold=True)
for r in rows: ws.append([r[c] for c in COLS])
fmt = {"market_cap_usd_m": "#,##0.000", "price": "0.0000", "shares_m": "#,##0.000", "operating_income_m": "#,##0.000;[Red]-#,##0.000",
       "tax_rate": "0.0%", "nopat_m": "#,##0.000;[Red]-#,##0.000"}
widths = {"company": 34, "sic_desc": 28, "description": 70, "nopat_sentence": 70, "flags": 28}
for i, c in enumerate(COLS, 1):
    L = get_column_letter(i)
    ws.column_dimensions[L].width = widths.get(c, 14)
    if c in fmt:
        for cell in ws[L][1:]: cell.number_format = fmt[c]
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
wb.save(OUT / f"{name}.xlsx")

# checks
caps = [r["market_cap_usd_m"] for r in rows]
def recompute(r):
    if r["operating_income_m"] is None: return r["nopat_m"] is None
    op = r["operating_income_m"]
    exp = op if op < 0 else op * (1 - (r["tax_rate"] if r["tax_rate"] is not None else 0.21))
    return abs(exp - r["nopat_m"]) <= max(1e-3, 0.002 * abs(exp))
chk = {
    "rows": len(rows),
    "all_caps_lt_1B": all(c < 1000 for c in caps),
    "ascending": all(a <= b for a, b in zip(caps, caps[1:])),
    "duplicate_ciks": len(rows) - len({r["cik"] for r in rows}),
    "duplicate_tickers": len(rows) - len({r["ticker"] for r in rows}),
    "price_date_max_age_days": max((today - date.fromisoformat(r["price_date"])).days for r in rows),
    "shares_date_gt150d_rows": sum((today - date.fromisoformat(r["shares_date"])).days > 150 for r in rows),
    "shares_date_gt150d_all_flagged": all("shares_stale" in r["flags"] for r in rows if (today - date.fromisoformat(r["shares_date"])).days > 150),
    "nopat_recompute_failures": [r["ticker"] for r in rows if not recompute(r)],
    "missing_description": sum(not r["description"] for r in rows),
    "cap_min_m": min(caps), "cap_median_m": statistics.median(caps), "cap_max_m": max(caps),
    "exchange_mix": dict(Counter(r["exchange"] for r in rows)),
    "nopat_positive": sum(1 for r in rows if r["nopat_m"] is not None and r["nopat_m"] > 0),
    "nopat_negative": sum(1 for r in rows if r["nopat_m"] is not None and r["nopat_m"] <= 0),
    "nopat_unavailable": sum(1 for r in rows if r["nopat_m"] is None),
    "description_sic_fallbacks": sum("desc_sic_fallback" in r["flags"] for r in rows),
    "flag_counts": dict(Counter(f for r in rows for f in r["flags"].split(";") if f)),
    "stage_stats": {"universe": json.loads((OUT / ("universe_test.json" if TEST else "universe.json")).read_text())["stats"],
                    "caps": CAPS["stats"]},
}
chk["PASS"] = (chk["all_caps_lt_1B"] and chk["ascending"] and chk["duplicate_ciks"] == 0
               and not chk["nopat_recompute_failures"] and chk["price_date_max_age_days"] <= 30
               and chk["shares_date_gt150d_all_flagged"])
(OUT / f"checks{sfx}.json").write_text(json.dumps(chk, indent=1))
print(json.dumps(chk, indent=1))

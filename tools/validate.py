"""Mechanical and data validation for one ticker.

Usage: python3 tools/validate.py TICKER
1. Recalculates TICKER_EPV.xlsx in LibreOffice and compares key rows with summary.json (Python replica).
2. Runs data tie outs and plausibility checks.
Writes output/TICKER/validation.json with status PASS, WARN, or FAIL.
"""
import argparse

import openpyxl

from common import YEAR_COLS, load_inputs, load_json, out_dir, save_json
from recalc import recalc

# summary.json key -> Financials row
COMPARE = {"revenue": 31, "operating_profit": 38, "net_income": 48, "ebitda": 80, "nopat": 123,
           "invested_capital": 130, "roic": 133, "epv": 140, "ev": 15}
CRITICAL = ("revenue", "operating_income", "pretax", "taxes", "net_income", "shares", "price",
            "ta", "tl", "tca", "tcl", "cash")


def close(a, b, rel=1e-6, abs_=1e-6):
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= max(abs_, rel * max(abs(a), abs(b)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ticker")
    t = ap.parse_args().ticker.upper()
    d = out_dir(t)
    inp, summ = load_inputs(t), load_json(d / "summary.json")
    years = [str(y) for y in inp["years"]]
    f = inp["fields"]
    fails, warns, info = [], [], []

    # 1. Workbook vs Python replica
    calc_path = d / f"{t}_EPV_recalc.xlsx"
    recalc(d / f"{t}_EPV.xlsx", calc_path)
    ws = openpyxl.load_workbook(calc_path, data_only=True)["Financials"]
    for col, y in zip(YEAR_COLS, years):
        for key, row in COMPARE.items():
            wb_v, py_v = ws[f"{col}{row}"].value, summ["by_year"][y][key]
            if isinstance(wb_v, str):
                wb_v = None
            if not close(wb_v, py_v, rel=1e-6):
                fails.append(f"{y} {key}: workbook {wb_v} vs python {py_v}")
    ps = summ["per_share_latest"] or {}
    for cell, key in [("D21", "invested_capital"), ("D22", "pv_comp_adv"), ("D23", "epv"),
                      ("D24", "growth_value"), ("D25", "ev")]:
        if not close(ws[cell].value, ps.get(key)):
            fails.append(f"summary {cell} {key}: workbook {ws[cell].value} vs python {ps.get(key)}")
    calc_path.unlink()

    # 2. Data completeness
    missing = {}
    for m in inp["missing"]:
        field, y = m.split(":")
        missing.setdefault(field, []).append(y)
    for field, ys in missing.items():
        msg = f"missing {field} for {', '.join(ys)}"
        if field == "cogs":
            msg += " (only affects gross margin row 34; operating profit still ties)"
        (fails if field in CRITICAL else warns).append(msg)

    # 3. Tie outs to reported figures
    by = summ["by_year"]
    for y in years:
        for key, ref in [("operating_profit", "operating_income"), ("net_income", "net_income")]:
            rep = inp["checks"][ref].get(y)
            if rep is not None and not close(by[y][key], rep, rel=0.005, abs_=0.5):
                fails.append(f"{y} {key} {by[y][key]:.1f} does not tie to reported {rep:.1f}")

    # 4. Plausibility
    for i, y in enumerate(years):
        b, v = by[y], (lambda k: f.get(k, {}).get(y) or 0.0)
        if v("cash") > v("tca") + 0.5 or v("tca") > v("ta") + 0.5 or v("tcl") > v("tl") + 0.5:
            fails.append(f"{y} balance sheet subtotals inconsistent (cash/TCA/TA or TCL/TL)")
        if v("cur_debt_lease") > v("tcl") + 0.5:
            fails.append(f"{y} current debt + lease {v('cur_debt_lease'):.0f} exceeds current liabilities "
                         f"{v('tcl'):.0f} (row 63 mapping)")
        if b["tax_rate"] is not None and not 0 <= b["tax_rate"] <= 0.40:
            warns.append(f"{y} tax rate {b['tax_rate']:.0%} distorts NOPAT (row 121 uses actual rate)")
        if b["invested_capital"] <= 0:
            warns.append(f"{y} invested capital {b['invested_capital']:.0f} <= 0, ROIC meaningless")
        elif b["roic"] is not None and abs(b["roic"]) > 1.0:
            warns.append(f"{y} ROIC {b['roic']:.0%} extreme (small or negative working capital base)")
        if abs(v("minority")) > 0.05 * max(abs(b["net_income"]), 1):
            warns.append(f"{y} row 46 plug {v('minority'):.0f} (NCI, equity income or discontinued ops)")
        if i:
            s0, s1 = inp["shares_current_basis"].get(years[i - 1]), inp["shares_current_basis"].get(y)
            if s0 and s1 and abs(s1 / s0 - 1) > 0.25:
                warns.append(f"{y} shares changed {s1 / s0 - 1:+.0%}; check split adjustment")
    for fld in ("ppe_gross", "accum_dep"):
        gaps = [y for y in years[:6] if f.get(fld, {}).get(y) is None]  # template cols D..I
        if gaps:
            info.append(f"{fld} null for {', '.join(gaps)} (template columns D..I); gross PP&E rows not populated")
    ic = {y: by[y]["invested_capital"] for y in years}
    if "2018" in ic and "2019" in ic and ic["2018"] > 0 and ic["2019"] / ic["2018"] > 1.3:
        info.append(f"IC jumps {ic['2019'] / ic['2018'] - 1:+.0%} in 2019: ASC 842 operating leases "
                    "enter row 57; ROIC before and after 2019 is not comparable")
    if summ.get("epv_alt_row100") and ps.get("epv"):
        info.append(f"Template has two EPVs: row 100 = {summ['epv_alt_row100']:.2f}/sh, "
                    f"row 140 (used in summary) = {ps['epv']:.2f}/sh")
    info += inp["notes"]

    status = "FAIL" if fails else "WARN" if warns else "PASS"
    save_json(d / "validation.json", {"ticker": t, "status": status, "fail": fails, "warn": warns, "info": info})
    print(f"{t}: {status} (fail {len(fails)}, warn {len(warns)}, info {len(info)})")


if __name__ == "__main__":
    main()
